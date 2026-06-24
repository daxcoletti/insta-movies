#!/usr/bin/env python3
"""
find_torrents.py — para cada película de la planilla busca el mejor candidato de
descarga por torrent y agrega el magnet link como columna "E" (antes de Caption).

⚠️ Uso dual: busca torrents de películas que pueden tener copyright. Es para uso
personal; verificá la legalidad en tu jurisdicción y preferí fuentes legales o de
dominio público cuando puedas.

Cómo funciona:
  1. Lee <PERFIL>_peliculas.csv (lo genera build_spreadsheet.py).
  2. Los títulos están en español. Resuelve el título ORIGINAL (idioma original,
     transliterado) y el título internacional/inglés con Claude (en lotes), porque
     los releases en los indexadores usan casi siempre el título original/inglés.
  3. Busca en apibay (API de The Pirate Bay) con "<título> <año>" y elige el mejor
     candidato: más seeders, exigiendo que coincida el año o las palabras del
     título (para no meter falsos positivos). Arma el magnet link.
  4. Reescribe la planilla con columna E = "Torrent" (Caption pasa a F),
     en CSV (UTF-8 BOM) y XLSX.

Uso:
    export ANTHROPIC_API_KEY=sk-ant-...
    python find_torrents.py [PERFIL]          # default: juan.amonda
"""

import csv
import json
import os
import re
import sys
import time
import unicodedata
import urllib.parse

import anthropic
import requests

MODEL = "claude-sonnet-4-6"
DEFAULT_PROFILE = "juan.amonda"
APIBAY = "https://apibay.org/q.php"
BATCH = 40
HTTP_HEADERS = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36"}

# Trackers públicos para el magnet (los que usa apibay por defecto).
TRACKERS = [
    "udp://tracker.opentrackr.org:1337/announce",
    "udp://open.stealth.si:80/announce",
    "udp://tracker.torrent.eu.org:451/announce",
    "udp://open.tracker.cl:1337/announce",
    "udp://exodus.desync.com:6969/announce",
    "udp://tracker.openbittorrent.com:6969/announce",
]

# Categorías de apibay que son películas (excluye TV: 205, 208).
MOVIE_CATS = {"201", "202", "204", "207", "209", "299"}
STOPWORDS = {"the", "a", "an", "de", "la", "el", "los", "las", "un", "una", "y",
             "of", "and", "to", "in"}


def die(msg, code=1):
    print(f"Error: {msg}", file=sys.stderr)
    sys.exit(code)


def norm(s: str) -> str:
    """minúsculas, sin acentos, sin puntuación -> para comparar."""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def clean_query(s: str) -> str:
    """Query para apibay: ascii sin acentos, apóstrofos eliminados (Winter's->Winters),
    resto de puntuación a espacio. apibay devuelve 0 resultados si la query trae
    apóstrofos u otra puntuación."""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("'", "").replace("’", "")  # apóstrofo recto y tipográfico
    s = re.sub(r"[^A-Za-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def queries_for(title: str, year: str) -> list:
    """Variantes de query para un título: completo, recortado al título principal
    (antes del ':' y sin un 'or/aka' colgando) y primeras 5 palabras. Cada una con y
    sin año. Recortar ayuda con subtítulos largos: 'Dr. Strangelove or: How I...' ->
    'Dr Strangelove', que es como se nombran los releases."""
    full = clean_query(title)
    head = re.split(r"\s*:\s*", title, maxsplit=1)[0]
    head = re.sub(r"\s+(or|aka|o)$", "", head, flags=re.IGNORECASE)
    head = clean_query(head)
    short = " ".join(full.split()[:5])
    out, seen = [], set()
    for base in (full, head, short):
        for q in ((f"{base} {year}", base) if year else (base,)):
            q = q.strip()
            k = q.lower()
            if q and k not in seen:
                seen.add(k)
                out.append(q)
    return out


# ---------- 1. Resolución de títulos originales con Claude ----------

RESOLVE_SYSTEM = (
    "Sos un experto en cine internacional. Te doy una lista de películas con su "
    "título en español y el año. Para cada una devolvé el título ORIGINAL (en su "
    "idioma original; si usa otro alfabeto, translitéralo al alfabeto latino) y el "
    "título internacional/inglés más usado en los releases. Respondé ÚNICAMENTE con "
    "un array JSON válido, sin texto ni explicación alrededor. Formato exacto: "
    '[{"i": <indice>, "original": "<titulo original>", "english": "<titulo ingles>"}]. '
    "Si no reconocés una película, repetí en ambos campos el título que te di."
)


def resolve_titles(client, movies):
    """movies: lista de (idx, titulo, anio). Devuelve {idx: (original, english)}."""
    out = {}
    for start in range(0, len(movies), BATCH):
        chunk = movies[start:start + BATCH]
        payload = [{"i": i, "titulo": t, "anio": y} for (i, t, y) in chunk]
        resp = client.messages.create(
            model=MODEL,
            max_tokens=4000,
            system=RESOLVE_SYSTEM,
            messages=[{"role": "user", "content": json.dumps(payload, ensure_ascii=False)}],
        )
        text = next((b.text for b in resp.content if b.type == "text"), "").strip()
        text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
        try:
            data = json.loads(text)
        except json.JSONDecodeError:
            m = re.search(r"\[.*\]", text, re.DOTALL)
            data = json.loads(m.group(0)) if m else []
        for row in data:
            i = row.get("i")
            if i is not None:
                out[i] = (row.get("original") or "", row.get("english") or "")
        print(f">> Títulos resueltos: {min(start + BATCH, len(movies))}/{len(movies)}",
              file=sys.stderr)
    return out


# ---------- 2. Búsqueda en apibay ----------

def apibay_search(query):
    try:
        r = requests.get(APIBAY, params={"q": query}, headers=HTTP_HEADERS, timeout=15)
        r.raise_for_status()
        data = r.json()
    except Exception as e:
        print(f"   ! apibay error ({query!r}): {e}", file=sys.stderr)
        return []
    results = []
    for it in data:
        h = it.get("info_hash", "")
        if not h or set(h) == {"0"}:  # sentinela "No results returned"
            continue
        results.append(it)
    return results


def acceptable(name_norm, title_words, year):
    """¿El candidato corresponde a la película buscada?"""
    year_ok = bool(year) and year in name_norm
    if not title_words:
        return year_ok
    matched = sum(1 for w in title_words if w in name_norm)
    ratio = matched / len(title_words)
    if year_ok and ratio >= 0.5:
        return True
    if matched >= 2 and ratio >= 0.8:
        return True
    return False


def quality_bonus(name_norm):
    if "2160p" in name_norm or "4k" in name_norm:
        return 3
    if "1080p" in name_norm:
        return 2
    if "720p" in name_norm:
        return 1
    return 0


def best_torrent(titles, year):
    """titles: lista de títulos a probar (en orden). Devuelve (magnet, info) o (None, None)."""
    seen_queries = set()
    for title in titles:
        if not title:
            continue
        title_words = [w for w in norm(title).split() if len(w) >= 3 and w not in STOPWORDS]
        for q in queries_for(title, year):
            qn = q.lower().strip()
            if not qn or qn in seen_queries:
                continue
            seen_queries.add(qn)
            time.sleep(0.4)  # cortesía con apibay
            best = None
            for it in apibay_search(q):
                if it.get("category", "") in {"205", "208"}:  # TV, no
                    continue
                if it.get("category", "") and it["category"] not in MOVIE_CATS \
                        and not it["category"].startswith("2"):
                    continue
                name_norm = norm(it.get("name", ""))
                if not acceptable(name_norm, title_words, year):
                    continue
                seeders = int(it.get("seeders", 0) or 0)
                score = seeders + quality_bonus(name_norm) * 50 \
                    + (200 if year and year in name_norm else 0)
                if best is None or score > best[0]:
                    best = (score, it, seeders)
            if best:
                _, it, seeders = best
                magnet = build_magnet(it["info_hash"], it.get("name", ""))
                info = {"name": it.get("name", ""), "seeders": seeders,
                        "matched_title": title, "query": q}
                return magnet, info
    return None, None


def build_magnet(info_hash, name):
    parts = [f"magnet:?xt=urn:btih:{info_hash.lower()}",
             "dn=" + urllib.parse.quote(name)]
    parts += ["tr=" + urllib.parse.quote(t) for t in TRACKERS]
    return "&".join(parts)


# ---------- 3. Leer / escribir la planilla ----------

def main():
    profile = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PROFILE
    csv_path = f"{profile}_peliculas.csv"
    if not os.path.isfile(csv_path):
        die(f"no existe {csv_path}. Generá la planilla primero:\n"
            f"  python build_spreadsheet.py {profile} --login <TU_USUARIO_IG>")
    if not os.environ.get("ANTHROPIC_API_KEY"):
        die("falta ANTHROPIC_API_KEY.  export ANTHROPIC_API_KEY=sk-ant-...")

    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    if not rows:
        die(f"{csv_path} está vacío.")

    client = anthropic.Anthropic()

    # Resolver títulos originales (una sola vez por título único).
    uniq = {}  # (titulo, anio) -> idx
    movies = []
    for i, r in enumerate(rows):
        key = (r["Película"], r["Año"])
        if key not in uniq:
            uniq[key] = len(movies)
            movies.append((len(movies), r["Película"], r["Año"]))
    print(f">> Resolviendo títulos originales de {len(movies)} películas con {MODEL} ...",
          file=sys.stderr)
    resolved = resolve_titles(client, movies)

    # Buscar torrents.
    print(f">> Buscando torrents en apibay para {len(rows)} posts ...", file=sys.stderr)
    cache = {}  # (titulo, anio) -> (magnet, info)
    hits = 0
    for n, r in enumerate(rows, 1):
        key = (r["Película"], r["Año"])
        if key not in cache:
            idx = uniq[key]
            orig, eng = resolved.get(idx, ("", ""))
            # orden de preferencia: inglés, original, español
            candidates = [eng, orig, r["Película"]]
            cache[key] = best_torrent(candidates, r["Año"])
        magnet, info = cache[key]
        r["Torrent"] = magnet or ""
        if magnet:
            hits += 1
            print(f"[{n}/{len(rows)}] {r['Película']} ({r['Año']}) -> "
                  f"{info['name'][:55]} [{info['seeders']}s]", file=sys.stderr)
        else:
            print(f"[{n}/{len(rows)}] {r['Película']} ({r['Año']}) -> sin torrent",
                  file=sys.stderr)

    # Escribir CSV con la nueva columna E (Torrent antes de Caption).
    cols = ["Fecha", "Película", "Año", "Link", "Torrent", "Caption"]
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})
    print(f">> CSV actualizado: {csv_path}", file=sys.stderr)

    # XLSX.
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Películas"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, r in enumerate(rows, start=2):
            ws.append([r.get(c, "") for c in cols])
            if r.get("Link"):
                cell = ws.cell(row=i, column=4)
                cell.hyperlink = r["Link"]
                cell.font = Font(color="0563C1", underline="single")
        widths = {"A": 12, "B": 32, "C": 7, "D": 40, "E": 60, "F": 80}
        for col, wdt in widths.items():
            ws.column_dimensions[col].width = wdt
        ws.freeze_panes = "A2"
        xlsx_path = f"{profile}_peliculas.xlsx"
        wb.save(xlsx_path)
        print(f">> XLSX actualizado: {xlsx_path}", file=sys.stderr)
    except ImportError:
        print(">> (openpyxl no instalado; salté el .xlsx)", file=sys.stderr)

    print(f">> Listo: {hits}/{len(rows)} posts con torrent encontrado.", file=sys.stderr)


if __name__ == "__main__":
    main()
